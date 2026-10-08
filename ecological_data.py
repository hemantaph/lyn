"""Read the supplied ecological workbooks without modifying their contents.

Excel coordinates below are explicit so the irregular worksheet layouts remain
auditable. Plotting and statistical choices are documented in the notebooks.
"""
from pathlib import Path
import re

import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
KEYS = ["Wetland", "Year", "Season", "Site"]
SEASONS = ["Pre-monsoon", "Monsoon", "Post-monsoon", "Winter"]
WETLANDS = ["Loktak", "Pumlen"]


def read_block(ws, header_row, site_column, variable):
    """Read two years, two lakes, five sites, four seasons; retain cell provenance."""
    records = []
    for offset in [0, 8]:
        for dc in [0, 5]:
            r, c = header_row + offset, site_column + dc
            match = re.fullmatch(r"(Loktak|Pumlen) (2022|2023)", str(ws.cell(r, c).value).strip())
            if not match:
                raise ValueError(f"Unexpected group header: {ws.title}!{ws.cell(r,c).coordinate}")
            assert [ws.cell(r+1,c+j).value for j in range(1,5)] == ["Pre","Mon","Post","Win"]
            for dr in range(2, 7):
                site = ws.cell(r+dr, c).value
                assert site in range(1,6)
                for j, season in enumerate(SEASONS, 1):
                    cell = ws.cell(r+dr, c+j)
                    if not isinstance(cell.value, (float, int)):
                        raise ValueError(f"Non-numeric measurement: {ws.title}!{cell.coordinate}")
                    records.append(dict(Wetland=match[1], Year=int(match[2]), Season=season,
                                        Site=int(site), **{variable:float(cell.value)},
                                        Source=f"Scatter Plot.xlsx / {ws.title}!{cell.coordinate}"))
    result = pd.DataFrame(records)
    assert len(result) == 80 and not result.duplicated(KEYS).any()
    return result


def load_scatter(root=ROOT):
    wb = load_workbook(root / "Scatter Plot.xlsx", data_only=True)
    specs = {
        "Soil_stock": ("ScatterPlot1",4,4),
        "Soil_Temp": ("ScatterPlot1",6,18),
        "Soil_pH": ("ScatterPlot1",24,18),
        "BD": ("ScatterPlot1",42,18),
        "Moisture": ("ScatterPlot1",60,18),
        "N": ("ScatterPlot1",78,18),
        "P": ("ScatterPlot1",96,18),
        "K": ("ScatterPlot1",114,18),
        "TBC_block_unverified": ("ScatterPlot2",4,36),
        "Phyto_carbon": ("ScatterPlot3",8,9),
        "Water_Temp": ("ScatterPlot3",26,9),
        "CO2": ("ScatterPlot3",44,9),
        "Hardness": ("ScatterPlot3",62,9),
    }
    blocks = {name:read_block(wb[sheet],row,col,name) for name,(sheet,row,col) in specs.items()}
    wide = None
    for name, frame in blocks.items():
        part = frame.drop(columns="Source")
        wide = part if wide is None else wide.merge(part,on=KEYS,validate="one_to_one")
    return wide, blocks


def load_pca(root=ROOT):
    """Keep original numeric identifiers for auditing; decode only in notebooks."""
    frames = pd.read_excel(root / "PCA.xlsx", sheet_name=None, engine="openpyxl")
    for frame in frames.values():
        assert len(frame) == 80 and not frame.duplicated(KEYS).any()
        assert frame.notna().all().all()
        assert np.isfinite(frame.to_numpy(dtype=float)).all()
    return frames


def decode_pca(frame):
    """Mapping checked against named water/soil blocks in notebook 01."""
    frame = frame.copy()
    frame["Wetland"] = frame.Wetland.map({1:"Loktak",2:"Pumlen"})
    frame["Year"] = frame.Year.map({1:2022,2:2023})
    frame["Season"] = frame.Season.map(dict(enumerate(SEASONS,1)))
    assert frame[KEYS].notna().all().all()
    return frame


def load_line_summary(root=ROOT):
    ws = load_workbook(root / "Scatter Plot.xlsx",data_only=True)["Line_Graph"]
    rows = []
    for variable, start in [("Soil_stock",3),("TBC",8),("Phyto_carbon",13)]:
        for i in range(4):
            r = start+i
            for c, season in enumerate(SEASONS,6):
                mean, spread = map(float, str(ws.cell(r,c).value).split("±"))
                rows.append(dict(Variable=variable, Wetland=WETLANDS[i//2],
                                 Year=int(ws.cell(r,5).value),Season=season,
                                 Mean=mean,Reported_spread=spread,
                                 Source=f"Scatter Plot.xlsx / Line_Graph!{ws.cell(r,c).coordinate}"))
    return pd.DataFrame(rows)


def number(value):
    """Parse exported r values, including Unicode minus signs and star suffixes."""
    if value is None:
        return np.nan
    return float(str(value).replace("−","-").replace("*", "").strip())


def load_heatmaps(root=ROOT):
    wb = load_workbook(root / "Heatmap.xlsx", data_only=True)
    result = {}
    for sheet, header, first_col, size in [
        ("heatmap1",3,4,9),("heatmap2",4,5,9),
        ("heatmap3",5,5,14),("heatmap4",4,5,14),
    ]:
        ws = wb[sheet]
        labels = [str(ws.cell(header,first_col+j).value).strip() for j in range(size)]
        matrices = {}
        for key, offset in [("r",0),("p",1),("n",2)]:
            values=[]
            for i,label in enumerate(labels):
                r=header+1+3*i
                assert str(ws.cell(r,first_col-2).value).strip() == label
                assert ws.cell(r,first_col-1).value == "Pearson Correlation"
                values.append([number(ws.cell(r+offset,first_col+j).value) for j in range(size)])
            matrices[key] = pd.DataFrame(values,index=labels,columns=labels)
        r=matrices["r"].to_numpy()
        assert np.isfinite(r).all() and np.all(np.abs(r)<=1)
        assert np.allclose(r,r.T,atol=0.001) and np.allclose(np.diag(r),1)
        assert (matrices["n"] == 80).all().all()
        result[sheet]=matrices
    return result
